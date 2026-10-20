
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def make_docx(blog):
    document = Document()
    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)

    heading = document.add_heading(blog["title"] or "Untitled Blog", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for line in (blog.get("content") or "").splitlines():
        text = line.strip()

        if not text:
            continue
        elif text.startswith("### "):
            document.add_heading(text[4:], level=3)
        elif text.startswith("## "):
            document.add_heading(text[3:], level=2)
        elif text.startswith("# "):
            document.add_heading(text[2:], level=1)
        elif text.startswith(("- ", "* ", "• ")):
            document.add_paragraph(text[2:], style="List Bullet")
        else:
            document.add_paragraph(text)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
