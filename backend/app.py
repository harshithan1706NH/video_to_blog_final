
import os
import sys
import uuid
import subprocess
from pathlib import Path
from io import BytesIO

import cloudinary
import cloudinary.uploader
from dotenv import load_dotenv
from flask import Flask, request, jsonify, session, send_file
from flask_cors import CORS
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash

from audio_chunker import chunk_audio
from db import get_db_connection
from blog_generator import generate_blog, save_blog


BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY")
CORS(app, supports_credentials=True)
app.config["MAX_CONTENT_LENGTH"] = 500 * 1024 * 1024

UPLOAD_DIR = BASE_DIR / "uploads"
AUDIO_DIR = BASE_DIR / "processing"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
AUDIO_DIR.mkdir(parents=True, exist_ok=True)

cloudinary.config(
    cloud_name=os.getenv("CLOUDINARY_CLOUD_NAME"),
    api_key=os.getenv("CLOUDINARY_API_KEY"),
    api_secret=os.getenv("CLOUDINARY_API_SECRET"),
)


def get_owned_blog(cursor, blog_id, user_id):
    cursor.execute(
        """
        SELECT blog_id, title, content, created_at
        FROM blog_contents
        WHERE blog_id = %s AND user_id = %s
        """,
        (blog_id, user_id),
    )
    row = cursor.fetchone()

    if not row:
        return None

    return {
        "blog_id": row[0],
        "title": row[1],
        "content": row[2],
        "created_at": row[3].isoformat() if row[3] else None,
    }


@app.route("/signup", methods=["POST"])
def signup():
    data = request.get_json(silent=True) or {}

    name = data.get("name", "").strip()
    email = data.get("email", "").strip().lower()
    password = data.get("password", "")

    if not name or not email or not password:
        return jsonify({
            "success": False,
            "error": "Name, email, and password are required.",
        }), 400

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            "SELECT user_id FROM users WHERE email = %s",
            (email,),
        )

        if cursor.fetchone():
            return jsonify({
                "success": False,
                "error": "An account with this email already exists.",
            }), 409

        cursor.execute(
            """
            INSERT INTO users (name, email, password_hash)
            VALUES (%s, %s, %s)
            RETURNING user_id
            """,
            (name, email, generate_password_hash(password)),
        )

        user_id = cursor.fetchone()[0]
        connection.commit()

        session.clear()
        session["user_id"] = user_id
        session["user_name"] = name
        session["user_email"] = email

        return jsonify({
            "success": True,
            "message": "Account created successfully.",
            "user": {
                "user_id": user_id,
                "name": name,
                "email": email,
            },
        }), 201

    except Exception as e:
        if connection:
            connection.rollback()

        print("Signup error:", e)

        return jsonify({
            "success": False,
            "error": "Unable to create account.",
        }), 500

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


@app.route("/login", methods=["POST"])
def login():
    data = request.get_json(silent=True) or {}

    email = data.get("email", "").strip().lower()
    password = data.get("password", "")

    if not email or not password:
        return jsonify({
            "success": False,
            "error": "Email and password are required.",
        }), 400

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT user_id, name, email, password_hash
            FROM users
            WHERE email = %s
            """,
            (email,),
        )

        user = cursor.fetchone()

        if not user or not check_password_hash(user[3], password):
            return jsonify({
                "success": False,
                "error": "Invalid email or password.",
            }), 401

        session.clear()
        session["user_id"] = user[0]
        session["user_name"] = user[1]
        session["user_email"] = user[2]

        return jsonify({
            "success": True,
            "message": "Login successful.",
            "user": {
                "user_id": user[0],
                "name": user[1],
                "email": user[2],
            },
        }), 200

    except Exception as e:
        print("Login error:", e)

        return jsonify({
            "success": False,
            "error": "Unable to log in.",
        }), 500

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


@app.route("/session", methods=["GET"])
def check_session():
    if "user_id" not in session:
        return jsonify({
            "success": True,
            "logged_in": False,
        }), 200

    return jsonify({
        "success": True,
        "logged_in": True,
        "user": {
            "user_id": session["user_id"],
            "name": session.get("user_name"),
            "email": session.get("user_email"),
        },
    }), 200


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()

    return jsonify({
        "success": True,
        "message": "Logged out successfully.",
    }), 200


@app.route("/upload", methods=["POST"])
def upload_video():
    if "user_id" not in session:
        return jsonify({
            "success": False,
            "error": "Please log in before uploading a video.",
        }), 401

    if not all([
        os.getenv("CLOUDINARY_CLOUD_NAME"),
        os.getenv("CLOUDINARY_API_KEY"),
        os.getenv("CLOUDINARY_API_SECRET"),
    ]):
        return jsonify({
            "success": False,
            "error": "Cloudinary configuration is missing.",
        }), 500

    if "video" not in request.files:
        return jsonify({
            "success": False,
            "error": "No video file was provided.",
        }), 400

    video = request.files["video"]

    if not video.filename:
        return jsonify({
            "success": False,
            "error": "Please select a video file.",
        }), 400

    if Path(video.filename).suffix.lower() != ".mp4":
        return jsonify({
            "success": False,
            "error": "Only MP4 videos are supported.",
        }), 400

    safe_filename = secure_filename(video.filename)
    local_video_path = UPLOAD_DIR / f"{uuid.uuid4().hex}_{safe_filename}"
    audio_path = None
    audio_chunks = []
    transcript_outputs = []
    connection = None
    cursor = None

    try:
        video.save(local_video_path)

        probe = subprocess.run(
            [
                "ffprobe",
                "-v", "error",
                "-show_entries", "format=duration",
                "-of", "default=noprint_wrappers=1:nokey=1",
                str(local_video_path),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=True,
        )

        duration_text = probe.stdout.strip()

        if not duration_text:
            return jsonify({
                "success": False,
                "error": "Unable to determine video duration.",
            }), 400

        duration_seconds = float(duration_text)

        if duration_seconds > 1800:
            return jsonify({
                "success": False,
                "error": "Video duration cannot exceed 30 minutes.",
            }), 400

        cloudinary_result = cloudinary.uploader.upload(
            str(local_video_path),
            resource_type="video",
            folder="video_to_blog",
        )

        video_url = cloudinary_result.get("secure_url")

        if not video_url:
            raise RuntimeError("Cloudinary did not return a video URL.")

        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            INSERT INTO videos (
                user_id, filename, video_url, duration_seconds
            )
            VALUES (%s, %s, %s, %s)
            RETURNING video_id
            """,
            (
                session["user_id"],
                safe_filename,
                video_url,
                duration_seconds,
            ),
        )

        video_id = cursor.fetchone()[0]
        connection.commit()

        audio_path = AUDIO_DIR / f"{uuid.uuid4().hex}.wav"

        subprocess.run(
            [
                "ffmpeg",
                "-y",
                "-i", str(local_video_path),
                "-vn",
                "-acodec", "pcm_s16le",
                "-ar", "16000",
                "-ac", "1",
                str(audio_path),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=True,
        )

        audio_chunks = chunk_audio(str(audio_path))

        if not audio_chunks:
            raise RuntimeError("Audio chunking did not produce any chunks.")

        transcript_parts = []

        for index, chunk in enumerate(audio_chunks):
            chunk_path = Path(chunk)
            transcript_path = AUDIO_DIR / f"{uuid.uuid4().hex}_transcript.txt"
            transcript_outputs.append(transcript_path)

            cursor.execute(
                """
                INSERT INTO audio_chunks (video_id, chunk_path, chunk_index)
                VALUES (%s, %s, %s)
                RETURNING chunk_id
                """,
                (video_id, str(chunk_path), index),
            )

            chunk_id = cursor.fetchone()[0]
            connection.commit()

            subprocess.run(
                [
                    sys.executable,
                    str(BASE_DIR / "parakeet_transcriber.py"),
                    str(chunk_path),
                    str(transcript_path),
                ],
                check=True,
            )

            transcript_text = ""

            if transcript_path.exists():
                transcript_text = transcript_path.read_text(
                    encoding="utf-8",
                    errors="replace",
                ).strip()

            cursor.execute(
                """
                INSERT INTO transcripts (video_id, chunk_id, transcript_text)
                VALUES (%s, %s, %s)
                """,
                (video_id, chunk_id, transcript_text),
            )

            connection.commit()

            if transcript_text:
                transcript_parts.append(transcript_text)

        combined_transcript = "\n".join(transcript_parts).strip()

        if not combined_transcript:
            raise RuntimeError("Transcription produced no text.")

        cursor.execute(
            "UPDATE videos SET transcript = %s WHERE video_id = %s",
            (combined_transcript, video_id),
        )
        connection.commit()

        title, content = generate_blog(combined_transcript)

        saved_blog = save_blog(
            video_id=video_id,
            user_id=session["user_id"],
            title=title,
            content=content,
        )

        return jsonify({
            "success": True,
            "message": "Video processed and blog generated successfully.",
            "video": {
                "video_id": video_id,
                "filename": safe_filename,
                "video_url": video_url,
                "duration_seconds": duration_seconds,
            },
            "blog": saved_blog,
        }), 201

    except subprocess.CalledProcessError as e:
        if connection:
            connection.rollback()

        print("Subprocess error:", e)

        if e.stderr:
            print(e.stderr)

        return jsonify({
            "success": False,
            "error": "Video processing failed. Check the backend terminal for details.",
        }), 500

    except Exception as e:
        if connection:
            connection.rollback()

        print("Upload error:", e)

        return jsonify({
            "success": False,
            "error": str(e),
        }), 500

    finally:
        if cursor:
            cursor.close()

        if connection:
            connection.close()

        for path in [local_video_path, audio_path, *transcript_outputs]:
            if path and Path(path).exists():
                try:
                    Path(path).unlink()
                except OSError:
                    pass

        for chunk in audio_chunks:
            chunk_path = Path(chunk)

            if chunk_path.exists():
                try:
                    chunk_path.unlink()
                except OSError:
                    pass


@app.route("/blogs", methods=["GET"])
def get_blogs():
    if "user_id" not in session:
        return jsonify({
            "success": False,
            "error": "Please log in to view your blogs.",
        }), 401

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT blog_id, title, content, created_at
            FROM blog_contents
            WHERE user_id = %s
            ORDER BY created_at DESC
            """,
            (session["user_id"],),
        )

        blogs = [
            {
                "blog_id": row[0],
                "title": row[1],
                "content": row[2],
                "created_at": row[3].isoformat() if row[3] else None,
            }
            for row in cursor.fetchall()
        ]

        return jsonify({
            "success": True,
            "blogs": blogs,
        }), 200

    except Exception as e:
        print("Get blogs error:", e)

        return jsonify({
            "success": False,
            "error": "Unable to retrieve blogs.",
        }), 500

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


@app.route("/blogs/<int:blog_id>", methods=["GET"])
def get_blog(blog_id):
    if "user_id" not in session:
        return jsonify({
            "success": False,
            "error": "Please log in to view this blog.",
        }), 401

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        blog = get_owned_blog(cursor, blog_id, session["user_id"])

        if not blog:
            return jsonify({
                "success": False,
                "error": "Blog not found.",
            }), 404

        return jsonify({
            "success": True,
            "blog": blog,
        }), 200

    except Exception as e:
        print("Get blog error:", e)

        return jsonify({
            "success": False,
            "error": "Unable to retrieve this blog.",
        }), 500

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


@app.route("/blogs/<int:blog_id>", methods=["PUT"])
def update_blog(blog_id):
    if "user_id" not in session:
        return jsonify({
            "success": False,
            "error": "Please log in to edit this blog.",
        }), 401

    data = request.get_json(silent=True) or {}
    title = data.get("title", "").strip()
    content = data.get("content", "").strip()

    if not title or not content:
        return jsonify({
            "success": False,
            "error": "Title and content are required.",
        }), 400

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            UPDATE blog_contents
            SET title = %s, content = %s
            WHERE blog_id = %s AND user_id = %s
            RETURNING blog_id, title, content, created_at
            """,
            (title, content, blog_id, session["user_id"]),
        )

        row = cursor.fetchone()

        if not row:
            connection.rollback()

            return jsonify({
                "success": False,
                "error": "Blog not found.",
            }), 404

        connection.commit()

        blog = {
            "blog_id": row[0],
            "title": row[1],
            "content": row[2],
            "created_at": row[3].isoformat() if row[3] else None,
        }

        return jsonify({
            "success": True,
            "message": "Blog updated successfully.",
            "blog": blog,
        }), 200

    except Exception as e:
        if connection:
            connection.rollback()

        print("Update blog error:", e)

        return jsonify({
            "success": False,
            "error": "Unable to update this blog.",
        }), 500

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


@app.route("/blogs/<int:blog_id>/export/docx", methods=["GET"])
def export_blog_docx(blog_id):
    if "user_id" not in session:
        return jsonify({
            "success": False,
            "error": "Please log in to export this blog.",
        }), 401

    connection = None
    cursor = None

    try:
        from docx import Document

        connection = get_db_connection()
        cursor = connection.cursor()

        blog = get_owned_blog(cursor, blog_id, session["user_id"])

        if not blog:
            return jsonify({
                "success": False,
                "error": "Blog not found.",
            }), 404

        document = Document()
        document.add_heading(blog["title"], level=0)

        for paragraph in blog["content"].splitlines():
            if paragraph.strip():
                document.add_paragraph(paragraph)

        output = BytesIO()
        document.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=f"blog_{blog_id}.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    except ImportError:
        return jsonify({
            "success": False,
            "error": "DOCX export requires python-docx. Install it with: pip install python-docx",
        }), 500

    except Exception as e:
        print("DOCX export error:", e)

        return jsonify({
            "success": False,
            "error": "Unable to export this blog as DOCX.",
        }), 500

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


@app.route("/blogs/<int:blog_id>/export/pdf", methods=["GET"])
def export_blog_pdf(blog_id):
    if "user_id" not in session:
        return jsonify({
            "success": False,
            "error": "Please log in to export this blog.",
        }), 401

    connection = None
    cursor = None

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        connection = get_db_connection()
        cursor = connection.cursor()

        blog = get_owned_blog(cursor, blog_id, session["user_id"])

        if not blog:
            return jsonify({
                "success": False,
                "error": "Blog not found.",
            }), 404

        from xml.sax.saxutils import escape

        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=inch,
            leftMargin=inch,
            topMargin=inch,
            bottomMargin=inch,
        )

        styles = getSampleStyleSheet()
        story = [
            Paragraph(escape(blog["title"]), styles["Title"]),
            Spacer(1, 12),
        ]

        for paragraph in blog["content"].splitlines():
            if paragraph.strip():
                story.append(
                    Paragraph(escape(paragraph), styles["BodyText"])
                )
                story.append(Spacer(1, 8))

        document.build(story)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=f"blog_{blog_id}.pdf",
            mimetype="application/pdf",
        )

    except ImportError:
        return jsonify({
            "success": False,
            "error": "PDF export requires reportlab. Install it with: pip install reportlab",
        }), 500

    except Exception as e:
        print("PDF export error:", e)

        return jsonify({
            "success": False,
            "error": "Unable to export this blog as PDF.",
        }), 500

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000)
