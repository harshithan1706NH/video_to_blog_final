
import os
from io import BytesIO
from xml.sax.saxutils import escape

from flask import Flask, jsonify, request, session, send_file
from flask_cors import CORS
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from db import get_db_connection

app = Flask(__name__)

# Must match your existing Flask app's secret key.
app.secret_key = os.getenv(
    "FLASK_SECRET_KEY",
    "development-secret-key"
)

origins = os.getenv(
    "EDITOR_ALLOWED_ORIGINS",
    "http://127.0.0.1:5500,http://localhost:5500"
).split(",")

CORS(
    app,
    origins=[origin.strip() for origin in origins if origin.strip()],
    supports_credentials=True
)


def get_blog(blog_id, user_id):
    connection = get_db_connection()

    try:
        cursor = connection.cursor()
        cursor.execute(
            """
            SELECT blog_id, user_id, title, content
            FROM blog_contents
            WHERE blog_id = %s AND user_id = %s
            """,
            (blog_id, user_id)
        )
        row = cursor.fetchone()
        cursor.close()

        if not row:
            return None

        return {
            "blog_id": row[0],
            "user_id": row[1],
            "title": row[2] or "",
            "content": row[3] or ""
        }
    finally:
        connection.close()


def logged_in_user():
    return session.get("user_id")


@app.get("/api/blogs/<int:blog_id>")
def load_blog(blog_id):
    user_id = logged_in_user()

    if user_id is None:
        return jsonify({"error": "Please log in first."}), 401

    try:
        blog = get_blog(blog_id, user_id)

        if blog is None:
            return jsonify({"error": "Blog not found."}), 404

        return jsonify({"blog": blog})
    except Exception:
        app.logger.exception("Loading blog failed")
        return jsonify({"error": "Could not load blog."}), 500


@app.put("/api/blogs/<int:blog_id>")
def update_blog(blog_id):
    user_id = logged_in_user()

    if user_id is None:
        return jsonify({"error": "Please log in first."}), 401

    data = request.get_json(silent=True)

    if not isinstance(data, dict):
        return jsonify({"error": "Invalid JSON data."}), 400

    title = data.get("title")
    content = data.get("content")

    if not isinstance(title, str) or not title.strip():
        return jsonify({"error": "Title is required."}), 400

    if not isinstance(content, str) or not content.strip():
        return jsonify({"error": "Content is required."}), 400

    connection = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            UPDATE blog_contents
            SET title = %s,
                content = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE blog_id = %s AND user_id = %s
            """,
            (title.strip(), content, blog_id, user_id)
        )

        if cursor.rowcount == 0:
            connection.rollback()
            cursor.close()
            return jsonify({"error": "Blog not found."}), 404

        connection.commit()
        cursor.close()

        return jsonify({
            "success": True,
            "message": "Blog saved successfully.",
            "blog": {
                "blog_id": blog_id,
                "title": title.strip(),
                "content": content
            }
        })

    except Exception:
        if connection:
            connection.rollback()

        app.logger.exception("Saving blog failed")
        return jsonify({"error": "Could not save blog."}), 500

    finally:
        if connection:
            connection.close()


@app.get("/api/blogs/<int:blog_id>/export/docx")
def export_word(blog_id):
    user_id = logged_in_user()

    if user_id is None:
        return jsonify({"error": "Please log in first."}), 401

    try:
        blog = get_blog(blog_id, user_id)

        if blog is None:
            return jsonify({"error": "Blog not found."}), 404

        document = Document()
        document.add_heading(blog["title"], level=1)

        for paragraph in blog["content"].splitlines():
            document.add_paragraph(paragraph)

        output = BytesIO()
        document.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=f"blog_{blog_id}.docx",
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        )

    except Exception:
        app.logger.exception("Word export failed")
        return jsonify({"error": "Could not export Word file."}), 500


@app.get("/api/blogs/<int:blog_id>/export/pdf")
def export_pdf(blog_id):
    user_id = logged_in_user()

    if user_id is None:
        return jsonify({"error": "Please log in first."}), 401

    try:
        blog = get_blog(blog_id, user_id)

        if blog is None:
            return jsonify({"error": "Blog not found."}), 404

        output = BytesIO()

        pdf = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=45,
            leftMargin=45,
            topMargin=45,
            bottomMargin=45
        )

        styles = getSampleStyleSheet()
        story = [
            Paragraph(escape(blog["title"]), styles["Title"]),
            Spacer(1, 16)
        ]

        for line in blog["content"].splitlines():
            if line.strip():
                story.append(
                    Paragraph(escape(line), styles["BodyText"])
                )
                story.append(Spacer(1, 8))
            else:
                story.append(Spacer(1, 8))

        pdf.build(story)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=f"blog_{blog_id}.pdf",
            mimetype="application/pdf"
        )

    except Exception:
        app.logger.exception("PDF export failed")
        return jsonify({"error": "Could not export PDF."}), 500


if __name__ == "__main__":
    app.run(
        host="127.0.0.1",
        port=5001,
        debug=False
    )
